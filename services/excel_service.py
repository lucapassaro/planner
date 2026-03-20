"""Import and export resource plans to/from Excel using openpyxl.

The extended Excel format (one row per person per activity) includes:
  Col A: Tipo (AM/EVO)
  Col B: Gruppo
  Col C: Team
  Col D: Attività
  Col E: Persona (resource name) — extension over the original format
  Col F: Stato
  Col G-R: Monthly allocations Jan-Dec (percentage %)
  Col S: Totale

Summary rows follow the data rows:
  - EVO subtotal (person-days)
  - AM subtotal
  - EVO+AM total
  - Giorni lavorativi per mese
  - FTE EVO, FTE AM, FTE TOT
"""

import io
from datetime import date
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db.engine import get_session
from db.models import Allocazione, Attivita, Piano, Risorsa
from services.allocazione_service import OverallocationError, create_allocazione
from services.attivita_service import create_attivita
from services.calendario_service import get_giorni_lavorativi, percentuale_to_giorni
from services.piano_service import create_piano, touch_piano
from services.risorsa_service import create_risorsa, get_all_risorse
from utils.app_logger import get_logger

_log = get_logger("excel")

# Header labels for the 12 months
MESI_IT = [
    "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
    "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre",
]

# Column indices (1-based) for openpyxl
COL_TIPO = 1
COL_GRUPPO = 2
COL_TEAM = 3
COL_NOME = 4
COL_PERSONA = 5
COL_STATO = 6
COL_MESI_START = 7   # January
COL_MESI_END = 18    # December
COL_TOTALE = 19


def export_to_excel(piano_id: int) -> bytes:
    """Export a plan to Excel bytes in the extended format.

    Args:
        piano_id: Primary key of the plan to export.

    Returns:
        Excel file content as bytes.

    Raises:
        ValueError: If the plan is not found.
    """
    with get_session() as session:
        piano = session.get(Piano, piano_id)
        if not piano:
            raise ValueError(f"Piano {piano_id} non trovato.")

        anno = piano.anno
        nome_piano = piano.nome

        # Load all activities and their allocations eagerly
        attivita_list = (
            session.query(Attivita)
            .filter(Attivita.piano_id == piano_id)
            .order_by(Attivita.tipo, Attivita.gruppo, Attivita.nome)
            .all()
        )

        # Build data structure: list of dicts (all values extracted in-session)
        rows_data = []
        for att in attivita_list:
            allocs_by_risorsa: Dict[str, Dict[int, float]] = {}
            for alloc in att.allocazioni:
                r_nome = alloc.risorsa.nome if alloc.risorsa else "N/A"
                parts = alloc.mese.split("-")
                m = int(parts[1])
                allocs_by_risorsa.setdefault(r_nome, {})[m] = alloc.percentuale

            att_data = {
                "tipo": att.tipo,
                "gruppo": att.gruppo,
                "team": att.team,
                "nome": att.nome,
                "stato": att.stato,
            }

            if allocs_by_risorsa:
                for r_nome, mesi_perc in sorted(allocs_by_risorsa.items()):
                    rows_data.append((att_data, r_nome, mesi_perc))
            else:
                rows_data.append((att_data, "", {}))

        # Pre-compute working days per month
        gg_lav = {m: get_giorni_lavorativi(anno, m) for m in range(1, 13)}

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scheduling"

    # === Header row ===
    _write_header(ws, anno)

    # === Data rows ===
    data_start_row = 2
    current_row = data_start_row

    evo_days: Dict[int, float] = {m: 0.0 for m in range(1, 13)}
    am_days: Dict[int, float] = {m: 0.0 for m in range(1, 13)}

    for att, r_nome, mesi_perc in rows_data:
        ws.cell(current_row, COL_TIPO).value = att["tipo"]
        ws.cell(current_row, COL_GRUPPO).value = att["gruppo"]
        ws.cell(current_row, COL_TEAM).value = att["team"]
        ws.cell(current_row, COL_NOME).value = att["nome"]
        ws.cell(current_row, COL_PERSONA).value = r_nome
        ws.cell(current_row, COL_STATO).value = att["stato"]

        row_total = 0.0
        for m in range(1, 13):
            perc = mesi_perc.get(m, 0.0)
            col = COL_MESI_START + (m - 1)
            if perc:
                ws.cell(current_row, col).value = round(perc, 1)
                days = percentuale_to_giorni(perc, anno, m)
                row_total += days
                if att["tipo"] == "EVO":
                    evo_days[m] += days
                else:
                    am_days[m] += days

        ws.cell(current_row, COL_TOTALE).value = round(row_total, 1)
        current_row += 1

    data_end_row = current_row - 1

    # === Summary rows ===
    _write_summary_rows(ws, current_row, anno, gg_lav, evo_days, am_days)

    # === Formatting ===
    _apply_formatting(ws, data_start_row, data_end_row, current_row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_from_excel(file_bytes: bytes, piano_nome: str) -> Dict:
    """Import a plan from Excel bytes (extended format or original format).

    Supports both:
    - Extended format (with Persona column at position 5)
    - Original format (no Persona column; imports activities only)

    Args:
        file_bytes: Raw Excel file bytes.
        piano_nome: Name for the new plan (must be unique).

    Returns:
        Dict with keys: 'piano_id', 'attivita_count', 'allocazioni_count',
                        'skipped_rows', 'warnings'.

    Raises:
        ValueError: If the plan name is taken or the file is malformed.
    """
    _log.info("=== Import Excel avviato: piano='%s' ===", piano_nome)

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    # Detect format by checking header row
    headers = [ws.cell(1, c).value for c in range(1, 8)]
    has_persona_col = _detect_persona_column(headers)
    _log.info("Formato rilevato: %s (righe dati: %d)",
              "esteso (con Persona)" if has_persona_col else "originale (senza Persona)",
              ws.max_row - 1)

    # Create the plan
    piano = create_piano(piano_nome)
    piano_id = piano.id
    _log.info("Piano creato: id=%d nome='%s'", piano_id, piano_nome)

    # Build resource cache (nome -> id)
    risorsa_cache: Dict[str, int] = {
        r.nome: r.id for r in get_all_risorse(only_active=False)
    }
    _log.debug("Risorse in cache pre-import: %d", len(risorsa_cache))

    attivita_cache: Dict[tuple, int] = {}  # (tipo, gruppo, team, nome) -> id
    attivita_count = 0
    allocazioni_count = 0
    skipped = 0
    warnings = []

    for row_idx in range(2, ws.max_row + 1):
        tipo_val = _cell_str(ws, row_idx, COL_TIPO)
        if not tipo_val or tipo_val not in ("AM", "EVO"):
            _log.debug("Riga %d ignorata: tipo='%s' (riga sommario o vuota)", row_idx, tipo_val)
            continue  # skip summary rows and blanks

        gruppo_val = _cell_str(ws, row_idx, COL_GRUPPO) or ""
        team_val = _cell_str(ws, row_idx, COL_TEAM) or ""
        nome_val = _cell_str(ws, row_idx, COL_NOME) or ""
        if not nome_val:
            _log.debug("Riga %d ignorata: nome attività vuoto", row_idx)
            skipped += 1
            continue

        if has_persona_col:
            persona_val = _cell_str(ws, row_idx, COL_PERSONA) or ""
            stato_val = _cell_str(ws, row_idx, COL_STATO) or "Da Confermare"
            mesi_col_offset = 0
        else:
            persona_val = ""
            stato_val = "Da Confermare"
            mesi_col_offset = -2  # shift month columns left (no Persona/Stato cols)

        # Validate stato
        valid_stati = ("Confermato", "Da Confermare", "Sospeso", "Chiuso")
        if stato_val not in valid_stati:
            _log.debug("Riga %d: stato '%s' non valido → sostituito con 'Da Confermare'",
                       row_idx, stato_val)
            stato_val = "Da Confermare"

        # Get or create activity
        att_key = (tipo_val, gruppo_val, team_val, nome_val)
        if att_key not in attivita_cache:
            try:
                att = create_attivita(
                    piano_id=piano_id,
                    gruppo=gruppo_val or "Generale",
                    tipo=tipo_val,
                    team=team_val or "N/A",
                    nome=nome_val,
                    stato=stato_val,
                )
                attivita_cache[att_key] = att.id
                attivita_count += 1
                _log.debug("Riga %d: attività creata id=%d [%s] '%s'",
                           row_idx, att.id, tipo_val, nome_val)
            except Exception as exc:
                msg = f"Riga {row_idx}: impossibile creare attività '{nome_val}' — {exc}"
                warnings.append(msg)
                _log.error(msg)
                skipped += 1
                continue

        attivita_id = attivita_cache[att_key]

        # Read monthly values
        if not persona_val:
            _log.debug("Riga %d: nessuna persona → allocazioni saltate", row_idx)
            continue  # no person to allocate to

        # Get or create resource
        if persona_val not in risorsa_cache:
            try:
                risorsa = create_risorsa(nome=persona_val, team=team_val or "N/A")
                risorsa_cache[persona_val] = risorsa.id
                _log.debug("Riga %d: risorsa creata id=%d '%s'",
                           row_idx, risorsa.id, persona_val)
            except ValueError:
                # Resource already exists (race condition) — fetch it
                with get_session() as session:
                    r = (
                        session.query(Risorsa)
                        .filter(Risorsa.nome == persona_val)
                        .first()
                    )
                    if r:
                        risorsa_cache[persona_val] = r.id
                        _log.debug("Riga %d: risorsa esistente id=%d '%s'",
                                   row_idx, r.id, persona_val)
                    else:
                        msg = f"Riga {row_idx}: risorsa '{persona_val}' non trovata e non creabile."
                        warnings.append(msg)
                        _log.warning(msg)
                        continue

        risorsa_id = risorsa_cache[persona_val]

        for m in range(1, 13):
            col = COL_MESI_START + (m - 1) + mesi_col_offset
            if col < 1:
                continue
            raw = ws.cell(row_idx, col).value
            if raw is None or raw == "":
                continue
            try:
                perc = float(raw)
            except (TypeError, ValueError):
                _log.debug("Riga %d mese %d: valore non numerico '%s' ignorato",
                           row_idx, m, raw)
                continue
            if perc <= 0:
                continue

            anno = int(piano_nome[-4:]) if piano_nome[-4:].isdigit() else 2026
            mese_str = f"{anno}-{m:02d}"

            try:
                create_allocazione(
                    attivita_id=attivita_id,
                    risorsa_id=risorsa_id,
                    mese=mese_str,
                    percentuale=perc,
                )
                allocazioni_count += 1
                _log.debug("Riga %d: allocazione creata att=%d ris='%s' mese=%s perc=%.1f%%",
                           row_idx, attivita_id, persona_val, mese_str, perc)
            except OverallocationError as exc:
                msg = f"Riga {row_idx}, mese {mese_str}: {exc}"
                warnings.append(msg)
                _log.warning(msg)
            except Exception as exc:
                msg = f"Riga {row_idx}, mese {mese_str}: errore imprevisto — {exc}"
                warnings.append(msg)
                _log.error(msg)

    touch_piano(piano_id)
    _log.info(
        "=== Import completato: %d attività, %d allocazioni, %d righe saltate, %d avvisi ===",
        attivita_count, allocazioni_count, skipped, len(warnings),
    )
    return {
        "piano_id": piano_id,
        "attivita_count": attivita_count,
        "allocazioni_count": allocazioni_count,
        "skipped_rows": skipped,
        "warnings": warnings,
    }


# ─── Private helpers ──────────────────────────────────────────────────────────


def _detect_persona_column(headers: List) -> bool:
    """Return True if the header row contains a Persona column."""
    for h in headers:
        if h and "persona" in str(h).lower():
            return True
    return False


def _cell_str(ws, row: int, col: int) -> Optional[str]:
    """Return cell value as stripped string or None."""
    val = ws.cell(row, col).value
    if val is None:
        return None
    return str(val).strip() or None


def _write_header(ws, anno: int) -> None:
    """Write the header row to the worksheet."""
    headers = [
        "Tipo", "Gruppo", "Team", "Attività", "Persona", "Stato"
    ] + MESI_IT + ["Totale (gg)"]

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="1F4E79")
    white = Font(bold=True, color="FFFFFF")

    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, label)
        cell.font = white
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    ws.row_dimensions[1].height = 30


def _write_summary_rows(ws, start_row: int, anno: int, gg_lav: Dict, evo_days: Dict, am_days: Dict) -> None:
    """Write summary/total rows below data rows."""
    labels_and_data = [
        ("EVO (gg)", evo_days),
        ("AM (gg)", am_days),
        ("TOTALE (gg)", {m: evo_days[m] + am_days[m] for m in range(1, 13)}),
        ("Giorni lavorativi", gg_lav),
    ]

    gray_fill = PatternFill("solid", fgColor="D9D9D9")
    bold = Font(bold=True)

    for i, (label, data) in enumerate(labels_and_data):
        r = start_row + i
        ws.cell(r, COL_NOME).value = label
        ws.cell(r, COL_NOME).font = bold

        row_total = 0.0
        for m in range(1, 13):
            val = round(data.get(m, 0.0), 1)
            ws.cell(r, COL_MESI_START + m - 1).value = val
            row_total += val

        ws.cell(r, COL_TOTALE).value = round(row_total, 1)

        for c in range(1, COL_TOTALE + 1):
            ws.cell(r, c).fill = gray_fill

    # FTE rows
    for i, (label, tipo) in enumerate([
        ("FTE EVO", "EVO"),
        ("FTE AM", "AM"),
        ("FTE TOT", "TOT"),
    ]):
        r = start_row + len(labels_and_data) + i
        ws.cell(r, COL_NOME).value = label
        ws.cell(r, COL_NOME).font = bold

        row_total = 0.0
        for m in range(1, 13):
            gg = gg_lav.get(m, 1)
            if tipo == "EVO":
                days = evo_days.get(m, 0.0)
            elif tipo == "AM":
                days = am_days.get(m, 0.0)
            else:
                days = evo_days.get(m, 0.0) + am_days.get(m, 0.0)

            fte = round(days / gg, 2) if gg else 0.0
            ws.cell(r, COL_MESI_START + m - 1).value = fte
            row_total += fte

        ws.cell(r, COL_TOTALE).value = round(row_total, 2)
        for c in range(1, COL_TOTALE + 1):
            ws.cell(r, c).fill = gray_fill


def _apply_formatting(ws, data_start: int, data_end: int, summary_start: int) -> None:
    """Apply column widths and data row formatting."""
    col_widths = {
        COL_TIPO: 6, COL_GRUPPO: 18, COL_TEAM: 8,
        COL_NOME: 35, COL_PERSONA: 18, COL_STATO: 14,
    }
    for m in range(13):
        col_widths[COL_MESI_START + m] = 9
    col_widths[COL_TOTALE] = 12

    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    # Freeze header
    ws.freeze_panes = "G2"

    # Alternate row fill for data rows
    alt_fill = PatternFill("solid", fgColor="EBF3FB")
    for r in range(data_start, data_end + 1):
        if r % 2 == 0:
            for c in range(1, COL_TOTALE + 1):
                if ws.cell(r, c).fill.fgColor.rgb == "00000000":
                    ws.cell(r, c).fill = alt_fill

    # Center month columns
    for r in range(data_start, summary_start + 10):
        for c in range(COL_MESI_START, COL_TOTALE + 1):
            ws.cell(r, c).alignment = Alignment(horizontal="center")
